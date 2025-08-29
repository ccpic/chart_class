#!/usr/bin/env python3
# -*- coding: utf-8 -*-

import pandas as pd
from ppt import PPT
from pptx.util import Cm, Pt


def test_add_table_with_index():
    """测试添加表格功能，包括显示索引"""

    # 创建测试数据，设置索引名称
    data = {
        "姓名": ["张三", "李四", "王五", "赵六"],
        "年龄": [25, 30, 35, 28],
        "部门": ["技术部", "市场部", "人事部", "财务部"],
        "薪资": [8000, 12000, 9000, 10000],
    }
    df = pd.DataFrame(data)
    df.index.name = "员工编号"  # 设置索引名称
    df.index = ["001", "002", "003", "004"]  # 设置索引值

    # 创建PPT对象
    try:
        ppt = PPT("template.pptx")
    except FileNotFoundError:
        print("未找到template.pptx文件，请确保存在模板文件")
        return

    # 添加内容页
    content = ppt.add_content_slide()

    # 设置标题
    content.set_title("员工信息表（显示索引）")

    # 添加表格 - 显示索引（默认）
    table1 = content.add_table(
        df=df,
        loc=content.body.top_left,
        anchor="top_left",
        width=Cm(16),  # 增加宽度以容纳索引列
        height=Cm(8),
        show_index=True,  # 显示索引
        header_bg_color="4472C4",  # 蓝色表头
        header_font_color="FFFFFF",  # 白色字体
        zebra_stripes=True,  # 启用斑马纹
    )

    # 添加表格 - 不显示索引
    table2 = content.add_table(
        df=df,
        loc=content.body.top_right,
        anchor="top_right",
        width=Cm(15),
        height=Cm(8),
        show_index=False,  # 不显示索引
        header_bg_color="70AD47",  # 绿色表头
        header_font_color="FFFFFF",
        zebra_stripes=True,
    )

    # 保存PPT
    ppt.save("test_table_with_index.pptx")
    print("测试完成！已生成 test_table_with_index.pptx 文件")


if __name__ == "__main__":
    test_add_table_with_index()
